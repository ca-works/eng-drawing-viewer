"""PowerPoint export: generated locally, editable titles and comments, good images."""
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches
from PySide6.QtGui import QColor, QImage

from eng_viewer.core import pptx_export


def _png(path: Path, w=1200, h=850, color=(255, 255, 255)):
    Image.new("RGB", (w, h), color).save(path, dpi=(220, 220))
    return str(path)


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text]


def test_new_presentation_layouts(tmp_path):
    img = _png(tmp_path / "a.png")
    out = tmp_path / "deck.pptx"
    spec = {
        "output": str(out), "mode": "new", "slide_size": "16:9", "title_mode": "title",
        "slides": [
            {"image": img, "img_w": 1200, "img_h": 850, "title": "Drawing 1.pdf", "marked": False},
            {"image": img, "img_w": 1200, "img_h": 850, "title": "Drawing 2.pdf", "marked": True,
             "comments": ["1. Chamfer edge 0.5 x 45°", "2. Check hole position ±0.05"]},
        ],
    }
    info = pptx_export.build(spec)
    assert info["slides"] == 2 and out.exists()
    prs = Presentation(str(out))
    assert prs.slide_width == Inches(13.333) and len(prs.slides) == 2
    s1, s2 = prs.slides
    assert s1.shapes.title is not None and s1.shapes.title.text == "Drawing 1.pdf"
    pics1 = [sh for sh in s1.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics1) == 1
    pic = pics1[0]
    assert abs((pic.left + pic.width / 2) - prs.slide_width / 2) < Inches(0.05)  # centred
    # Marked slide: picture on the right, editable numbered comments on the left.
    pic2 = [sh for sh in s2.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    boxes = [sh for sh in s2.shapes if sh.has_text_frame and "Chamfer" in sh.text_frame.text]
    assert boxes, _texts(s2)
    box = boxes[0]
    assert box.left < pic2.left
    paras = [p.text for p in box.text_frame.paragraphs]
    assert paras == ["1. Chamfer edge 0.5 x 45°", "2. Check hole position ±0.05"]
    assert box.text_frame.paragraphs[0].runs[0].font.bold  # bold number
    assert box.text_frame.word_wrap


def test_subtitle_and_none_modes(tmp_path):
    img = _png(tmp_path / "a.png")
    for mode in ("subtitle", "none"):
        out = tmp_path / f"{mode}.pptx"
        pptx_export.build({"output": str(out), "mode": "new", "title_mode": mode, "slide_size": "4:3",
                           "slides": [{"image": img, "img_w": 1200, "img_h": 850, "title": "X.pdf"}]})
        prs = Presentation(str(out))
        assert prs.slide_width == Inches(10)
        texts = _texts(prs.slides[0])
        assert (texts == ["X.pdf"]) if mode == "subtitle" else (texts == [])
        # No leftover empty placeholders ("Click to add title").
        assert all(ph.has_text_frame and ph.text_frame.text for ph in prs.slides[0].placeholders)


def _existing(tmp_path, n=3) -> Path:
    prs = Presentation()
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Existing {i + 1}"
    p = tmp_path / "existing.pptx"
    prs.save(p)
    return p


def test_insert_into_existing_at_position(tmp_path):
    p = _existing(tmp_path)
    img = _png(tmp_path / "a.png")
    pptx_export.build({"output": str(p), "template": str(p), "mode": "existing", "existing_mode": "insert",
                       "start_slide": 2, "title_mode": "title",
                       "slides": [{"image": img, "img_w": 10, "img_h": 10, "title": "New A"},
                                  {"image": img, "img_w": 10, "img_h": 10, "title": "New B"}]})
    titles = [s.shapes.title.text if s.shapes.title is not None else "" for s in Presentation(str(p)).slides]
    assert titles == ["Existing 1", "New A", "New B", "Existing 2", "Existing 3"]


def test_reuse_existing_slides_then_append(tmp_path):
    p = _existing(tmp_path, n=2)
    img = _png(tmp_path / "a.png")
    pptx_export.build({"output": str(p), "template": str(p), "mode": "existing", "existing_mode": "reuse",
                       "start_slide": 2, "title_mode": "subtitle",
                       "slides": [{"image": img, "img_w": 10, "img_h": 10, "title": "A.pdf"},
                                  {"image": img, "img_w": 10, "img_h": 10, "title": "B.pdf"}]})
    prs = Presentation(str(p))
    assert len(prs.slides) == 3  # slide 2 reused, one slide appended
    second = prs.slides[1]
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in second.shapes)
    assert "A.pdf" in _texts(second)


def test_export_controller_end_to_end(qtbot, render, corpus, tmp_path):
    """Queue entries → worker renders (with DPI) → deck; a MARKed page included."""
    from eng_viewer.ui.export_dialog import ExportController
    from eng_viewer.ui.queue_panel import QueueEntry

    marked = QImage(1600, 1130, QImage.Format.Format_RGB888)
    marked.fill(QColor("white"))
    entries = [
        QueueEntry(str(corpus["a1"]), 0),
        QueueEntry(str(corpus["multi"]), 3),
        QueueEntry(str(corpus["unicode"]), 0, marked_image=marked, comments="1. First\n2. Second", comment_count=2),
    ]
    out = tmp_path / "Drawings ünïcode.pptx"
    ctl = ExportController(render, entries, {"mode": "new", "output": str(out), "title_mode": "title", "quality": "standard",
                                             "slide_size": "16:9"})
    box = {}
    ctl.finished.connect(lambda info: box.setdefault("ok", info))
    ctl.failed.connect(lambda msg: box.setdefault("err", msg))
    ctl.start()
    qtbot.waitUntil(lambda: bool(box), timeout=90000)
    assert "ok" in box, box.get("err")
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Drawing 1.pdf", "Gearbox set (5 sheets).pdf", "Çhâssis — rév B (final) #12 [ä].pdf"]
    pic = [sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    w_px, h_px = pic.image.size
    assert abs(w_px - 2383.9 * 150 / 72) <= 2  # 150 DPI image of the A1 sheet
    assert pic.image.dpi[0] == 150
    assert "2. Second" in [p.text for sh in prs.slides[2].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
    assert not Path(ctl.tmp).exists()  # temporary images cleaned up
