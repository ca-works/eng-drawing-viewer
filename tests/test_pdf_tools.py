"""PDF Tools engine: every tool writes a correct new file and never touches inputs."""
import os
import sys
from pathlib import Path

import pymupdf
import pytest
from PIL import Image

from eng_viewer.core import pdf_tools as T
from eng_viewer.core.documents import is_document_name, open_document


@pytest.fixture()
def out(tmp_path):
    return tmp_path


def pages(path) -> int:
    with pymupdf.open(path) as d:
        return d.page_count


def test_parse_ranges():
    assert T.parse_ranges("", 4) == [0, 1, 2, 3]
    assert T.parse_ranges("1-3, 5", 6) == [0, 1, 2, 4]
    assert T.parse_ranges("3-end", 5) == [2, 3, 4]
    assert T.parse_ranges("4-2", 5) == [3, 2, 1]
    assert T.parse_ranges("odd", 5) == [0, 2, 4]
    assert T.parse_ranges("last", 7) == [6]
    with pytest.raises(T.ToolError):
        T.parse_ranges("9", 3)
    with pytest.raises(T.ToolError):
        T.parse_ranges("x", 3)


def test_document_names():
    assert is_document_name("a.PDF") and is_document_name("scan.HEIC") and is_document_name("p.tif")
    assert not is_document_name("photo.jpg", include_images=False)
    assert not is_document_name("._a.pdf") and not is_document_name("notes.txt")


def test_outputs_never_replace_inputs(corpus):
    src = str(corpus["multi"])
    with pytest.raises(T.ToolError, match="replace an input"):
        T.compress(src, src)


def test_organize_reorders_rotates_mixes_images_and_blanks(corpus, out, tmp_path):
    img = tmp_path / "photo.png"
    Image.new("RGB", (800, 600), "white").save(img)
    target = out / "organised.pdf"
    items = [
        {"path": str(corpus["multi"]), "page": 4},
        {"path": str(corpus["multi"]), "page": 0, "rotate": 90},
        {"image": str(img)},
        {"blank": [595, 842]},
        {"path": str(corpus["rotated"]), "page": 0},
    ]
    res = T.organize(items, str(target))
    assert res["outputs"] == [str(target)]
    with pymupdf.open(target) as d:
        assert d.page_count == 5
        assert "SHEET 5" in d[0].get_text()
        assert d[1].rotation == 90
        assert d[3].get_text().strip() == ""
        assert d[4].rotation == 90  # the source page's own rotation is kept


def test_merge_split_extract_remove_rotate(corpus, out):
    merged = out / "merged.pdf"
    T.merge([str(corpus["a1"]), str(corpus["multi"])], str(merged))
    assert pages(merged) == 6
    with pymupdf.open(merged) as d:
        assert [t[1] for t in d.get_toc()] == ["Drawing 1", "Gearbox set (5 sheets)"]
    res = T.split(str(corpus["multi"]), str(out / "split"), mode="each")
    assert len(res["outputs"]) == 5 and all(pages(p) == 1 for p in res["outputs"])
    res = T.split(str(corpus["multi"]), str(out / "split2"), mode="ranges", ranges="1-2; 3-5")
    assert [pages(p) for p in res["outputs"]] == [2, 3]
    T.select_pages(str(corpus["multi"]), str(out / "x.pdf"), "2,4", keep=True)
    assert pages(out / "x.pdf") == 2
    T.select_pages(str(corpus["multi"]), str(out / "y.pdf"), "2,4", keep=False)
    assert pages(out / "y.pdf") == 3
    T.rotate(str(corpus["multi"]), str(out / "r.pdf"), 180, "1")
    with pymupdf.open(out / "r.pdf") as d:
        assert [p.rotation for p in d] == [180, 0, 0, 0, 0]


def test_images_to_pdf_and_back(out, tmp_path):
    imgs = []
    for i, size in enumerate([(1200, 800), (600, 900)]):
        p = tmp_path / f"img{i}.jpg"
        Image.new("RGB", size, (200, 30, 30)).save(p, dpi=(150, 150))
        imgs.append(str(p))
    T.images_to_pdf(imgs, str(out / "imgs.pdf"), page_size="A4", landscape="auto", margin_mm=10)
    with pymupdf.open(out / "imgs.pdf") as d:
        assert d.page_count == 2
        assert d[0].rect.width > d[0].rect.height  # landscape image → landscape A4
        assert abs(d[1].rect.width - 595.3) < 1
    res = T.to_images(str(out / "imgs.pdf"), str(out / "png"), fmt="png", dpi=100)
    assert len(res["outputs"]) == 2
    with Image.open(res["outputs"][0]) as im:
        assert round(im.info["dpi"][0]) == 100
    res = T.to_images(str(out / "imgs.pdf"), str(out / "tif"), fmt="tiff", dpi=72, ranges="2")
    assert res["outputs"][0].endswith(".tiff")


def test_text_word_excel_powerpoint(corpus, out):
    res = T.to_text(str(corpus["a1"]), str(out / "a.md"), markdown=True)
    body = Path(res["outputs"][0]).read_text()
    assert "## Page 1" in body and "BRACKET ASSEMBLY" in body
    T.to_word(str(corpus["a4"]), str(out / "a.docx"), mode="both")
    from docx import Document

    doc = Document(out / "a.docx")
    assert any("SPEC SHEET" in p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes) == 1
    res = T.to_excel(str(corpus["a1"]), str(out / "a.xlsx"))
    assert Path(res["outputs"][0]).stat().st_size > 1000
    res = T.to_powerpoint(str(corpus["multi"]), str(out / "a.pptx"), dpi=72)
    from pptx import Presentation

    assert len(Presentation(out / "a.pptx").slides) == 5


def test_compress_repair_flatten(corpus, out, tmp_path):
    photo = tmp_path / "big.png"
    Image.effect_noise((2400, 1600), 60).convert("RGB").save(photo)
    T.images_to_pdf([str(photo)], str(tmp_path / "heavy.pdf"))
    res = T.compress(str(tmp_path / "heavy.pdf"), str(out / "small.pdf"), level="strong")
    assert res["after"] < res["before"]
    res = T.repair(str(corpus["a1"]), str(out / "fixed.pdf"))
    assert pages(out / "fixed.pdf") == 1
    with pytest.raises(T.ToolError):
        T.repair(str(corpus["root"] / "Problems" / "not really a pdf.pdf"), str(out / "nope.pdf"))
    src = tmp_path / "annot.pdf"
    d = pymupdf.open(str(corpus["a4"]))
    d[0].add_text_annot((100, 100), "review note")
    d.save(src)
    d.close()
    T.flatten(str(src), str(out / "flat.pdf"))
    with pymupdf.open(out / "flat.pdf") as f:
        assert list(f[0].annots()) == []


def test_numbers_watermark_stamp_overlay(corpus, out):
    T.page_numbers(str(corpus["multi"]), str(out / "n.pdf"), position="bottom-right", fmt="Page {n} of {total}")
    with pymupdf.open(out / "n.pdf") as d:
        assert "Page 3 of 5" in d[2].get_text()
    T.page_numbers(str(corpus["rotated"]), str(out / "nr.pdf"), position="bottom-center", fmt="PG-{n}-X")
    with pymupdf.open(out / "nr.pdf") as d:
        page = d[0]
        hit = page.search_for("PG-1-X")[0] * page.rotation_matrix  # search results are in unrotated space
        assert hit.y0 > page.rect.height * 0.9  # bottom of the page as displayed
        assert abs((hit.x0 + hit.x1) / 2 - page.rect.width / 2) < 30
    T.watermark(str(corpus["a4"]), str(out / "w.pdf"), text="DRAFT", style="diagonal")
    with pymupdf.open(out / "w.pdf") as d:
        assert "DRAFT" in d[0].get_text()
    T.watermark(str(corpus["a4"]), str(out / "s.pdf"), text="APPROVED", style="stamp", add_date=True)
    with pymupdf.open(out / "s.pdf") as d:
        assert "APPROVED" in d[0].get_text()
    T.overlay(str(corpus["multi"]), str(out / "s.pdf"), str(out / "o.pdf"))
    with pymupdf.open(out / "o.pdf") as d:
        assert all("APPROVED" in p.get_text() for p in d)


def test_crop_resize_tile_nup(corpus, out):
    T.crop(str(corpus["a4"]), str(out / "c.pdf"), mode="auto", margin_mm=2)
    with pymupdf.open(out / "c.pdf") as d, pymupdf.open(corpus["a4"]) as s:
        assert d[0].rect.width <= s[0].rect.width and d[0].rect.height <= s[0].rect.height
    T.crop(str(corpus["a4"]), str(out / "c2.pdf"), mode="margins", top=10, right=10, bottom=10, left=10)
    with pymupdf.open(out / "c2.pdf") as d:
        assert abs(d[0].rect.width - (595.3 - 20 * T.MM)) < 1
    T.resize(str(corpus["a1"]), str(out / "a3.pdf"), paper="A3")
    with pymupdf.open(out / "a3.pdf") as d:
        assert abs(d[0].rect.width - 420 * T.MM) < 1 and "BRACKET ASSEMBLY" in d[0].get_text()  # still vector text
    res = T.tile(str(corpus["a1"]), str(out / "t.pdf"), paper="A4", scale_pct=100, overlap_mm=10)
    assert pages(out / "t.pdf") >= 8  # A1 at 100 % needs at least 8 A4 sheets
    T.nup(str(corpus["multi"]), str(out / "u.pdf"), per_sheet=4, paper="A3")
    assert pages(out / "u.pdf") == 2


def test_metadata_protect_unlock(corpus, out):
    T.metadata(str(corpus["a4"]), str(out / "m.pdf"), title="Spec", author="Eng")
    with pymupdf.open(out / "m.pdf") as d:
        assert d.metadata["title"] == "Spec" and d.metadata["author"] == "Eng"
    T.metadata(str(out / "m.pdf"), str(out / "m2.pdf"), strip=True)
    with pymupdf.open(out / "m2.pdf") as d:
        assert not d.metadata.get("title")
    T.protect(str(corpus["a4"]), str(out / "p.pdf"), user_password="s3cret")
    with pymupdf.open(out / "p.pdf") as d:
        assert d.needs_pass
    with pytest.raises(T.ToolError, match="wrong"):
        T.unlock(str(out / "p.pdf"), str(out / "u.pdf"), password="nope")
    T.unlock(str(out / "p.pdf"), str(out / "u.pdf"), password="s3cret")
    with pymupdf.open(out / "u.pdf") as d:
        assert not d.needs_pass and "SPEC SHEET" in d[0].get_text()


def test_redaction_really_removes_text(corpus, out):
    res = T.redact(str(corpus["a1"]), str(out / "r.pdf"), terms="ENG-1000, AlMgSi1")
    assert "Removed" in res["summary"]
    with pymupdf.open(out / "r.pdf") as d:
        text = d[0].get_text()
        assert "ENG-1000" not in text and "AlMgSi1" not in text
        assert "BRACKET ASSEMBLY" in text  # everything else stays


@pytest.mark.skipif(T.ocr_engine() is None, reason="no OCR engine on this machine")
def test_ocr_makes_scans_searchable(out, tmp_path):
    # A "scan": text rendered to a bitmap, then wrapped in a PDF without a text layer.
    src = pymupdf.open()
    page = src.new_page(width=595, height=842)
    page.insert_text((72, 120), "TORQUE TO 25 NM", fontsize=28)
    pix = page.get_pixmap(dpi=200)
    scan = tmp_path / "scan.png"
    pix.save(scan)
    T.images_to_pdf([str(scan)], str(tmp_path / "scan.pdf"))
    res = T.ocr(str(tmp_path / "scan.pdf"), str(out / "ocr.pdf"))
    with pymupdf.open(out / "ocr.pdf") as d:
        assert "TORQUE" in d[0].get_text().upper()
        assert d[0].search_for("TORQUE")
    assert "searchable" in res["summary"]


def test_mark_pdf_keeps_vector_page_and_adds_overlay(corpus, out, tmp_path):
    ov = pymupdf.open()
    p = ov.new_page(width=1190.6, height=841.9)
    p.draw_circle((300, 300), 40, color=(1, 0, 0), width=3)
    p.insert_text((280, 310), "1", fontsize=30)
    ov.save(tmp_path / "ov.pdf")
    T.mark_pdf(str(corpus["multi"]), 2, str(tmp_path / "ov.pdf"), str(out / "marked.pdf"))
    with pymupdf.open(out / "marked.pdf") as d:
        assert d.page_count == 1
        text = d[0].get_text()
        assert "SHEET 3" in text and "1" in text
        assert d[0].get_drawings()  # vector content survived


@pytest.mark.parametrize("rot", [0, 90, 180, 270])
def test_mark_pdf_overlay_lands_where_it_was_drawn(rot, out, tmp_path):
    """MARK draws in displayed coordinates; the overlay must match on /Rotate pages too."""
    src = pymupdf.open()
    src.new_page(width=1190.6, height=841.9).set_rotation(rot)
    src.save(tmp_path / "src.pdf")
    w, h = src[0].rect.width, src[0].rect.height  # as displayed
    src.close()
    ov = pymupdf.open()
    op = ov.new_page(width=w, height=h)
    op.draw_rect(pymupdf.Rect(150, 110, 170, 130), color=(1, 0, 0), fill=(1, 0, 0))
    op.draw_rect(pymupdf.Rect(w - 60, h - 30, w - 20, h - 10), color=(0, 0, 1), fill=(0, 0, 1))
    ov.save(tmp_path / "ov.pdf")
    ov.close()
    T.mark_pdf(str(tmp_path / "src.pdf"), 0, str(tmp_path / "ov.pdf"), str(out / f"marked {rot}.pdf"))
    with pymupdf.open(out / f"marked {rot}.pdf") as d:
        assert d[0].rotation == rot and abs(d[0].rect.width - w) < 0.5
        pix = d[0].get_pixmap(alpha=False)
    assert pix.pixel(160, 120)[:3] == (255, 0, 0)
    assert pix.pixel(int(w - 40), int(h - 20))[:3] == (0, 0, 255)


def test_text_changes_between_revisions(corpus):
    changes = T.text_changes(str(corpus["a1"]), str(corpus["a1_2"]))
    removed = [l for s, l in changes if s == "-"]
    added = [l for s, l in changes if s == "+"]
    assert any("BRACKET" in l for l in removed) and any("HOUSING" in l for l in added)


def test_inputs_untouched(corpus, corpus_hashes):
    # Spot-check that the tools above did not write next to or into the corpus.
    import hashlib

    for path, (digest, _m) in corpus_hashes.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
