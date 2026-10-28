"""PDF editing and conversion tools (run inside render workers; no Qt).

Every tool reads its inputs and writes *new* files: an input is never modified
and an output may never replace an input. Vector content stays vector wherever
the operation allows it (organise, merge, resize, tile, N-up, stamps, page
numbers...). Each public function takes plain arguments and returns a dict with
``outputs`` (paths written) and an optional human-readable ``summary``.
"""
from __future__ import annotations

from contextlib import contextmanager
import datetime as _dt
import difflib
import io
import os
import re
import shutil
import secrets
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf

from .documents import extension, open_document
from .raster_limits import check_pixels, max_dpi_for, render_page
from .tool_support import find_office, unique_path
from .processes import run_conversion

MM = 72.0 / 25.4
PAPER_SIZES_MM = {
    "A0": (841, 1189), "A1": (594, 841), "A2": (420, 594), "A3": (297, 420), "A4": (210, 297), "A5": (148, 210),
    "Letter": (215.9, 279.4), "Legal": (215.9, 355.6), "Tabloid": (279.4, 431.8),
    "ANSI C": (431.8, 558.8), "ANSI D": (558.8, 863.6), "ANSI E": (863.6, 1117.6),
}


class ToolError(Exception):
    """A problem the user can understand and act on."""


# ---------------------------------------------------------------- helpers
def paper_size(name: str, landscape: bool = False) -> tuple[float, float]:
    w, h = PAPER_SIZES_MM[name]
    w, h = w * MM, h * MM
    return (h, w) if landscape else (w, h)


def parse_ranges(spec: str | None, count: int) -> list[int]:
    """'1-3, 5, 8-' → zero-based page indices (in the order given).

    Also accepts 'all', 'odd', 'even', 'last' and 'end' (e.g. '3-end').
    An empty spec means all pages.
    """
    spec = (spec or "").strip().lower()
    if spec in ("", "all", "*"):
        return list(range(count))
    if spec == "odd":
        return list(range(0, count, 2))
    if spec == "even":
        return list(range(1, count, 2))
    out: list[int] = []

    def num(tok: str, default: int) -> int:
        tok = tok.strip()
        if tok == "":
            return default
        if tok in ("last", "end", "z"):
            return count
        if not tok.isdigit():
            raise ToolError(f"“{tok}” is not a page number.")
        return int(tok)

    for part in re.split(r"[,;\s]+", spec):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start, end = num(a, 1), num(b, count)
            step = 1 if end >= start else -1
            rng = range(start, end + step, step)
        else:
            rng = [num(part, 1)]
        for n in rng:
            if not 1 <= n <= count:
                raise ToolError(f"Page {n} does not exist (the document has {count} page{'s' if count != 1 else ''}).")
            out.append(n - 1)
    if not out:
        raise ToolError("No pages selected.")
    return out


def _check_output(output: str, inputs: list[str]) -> Path:
    out = Path(output).expanduser()
    for src in inputs:
        try:
            if out.resolve() == Path(src).resolve() or (out.exists() and Path(src).exists() and out.samefile(src)):
                raise ToolError("The output would replace an input file. Choose a different name — originals are never overwritten.")
        except OSError:
            pass
    out.parent.mkdir(parents=True, exist_ok=True)
    return out


def _open(path: str, password: str | None = None):
    doc = None
    try:
        doc = open_document(pymupdf, path)
    except PermissionError:
        doc = pymupdf.open(path)
        if not password or not doc.authenticate(password):
            doc.close()
            raise ToolError(f"“{Path(path).name}” is password protected." + (" The password is wrong." if password else ""))
    except (ValueError, OSError) as exc:
        raise ToolError(f"“{Path(path).name}”: {exc}") from None
    return doc


def _as_pdf(path: str):
    """Open any supported document as a PDF document (images are converted)."""
    doc = _open(path)
    if doc.is_pdf:
        return doc
    pdf = pymupdf.open("pdf", doc.convert_to_pdf())
    doc.close()
    return pdf


@contextmanager
def _temporary_output(out: Path, suffix=".tmp"):
    fd, name = tempfile.mkstemp(dir=out.parent, prefix=f".{out.name}.", suffix=suffix)
    os.close(fd)
    tmp = Path(name)
    try:
        yield tmp
        os.replace(tmp, out)
    finally:
        tmp.unlink(missing_ok=True)


def _save(doc, out: Path, **kw) -> Path:
    """Save via a temporary file in the target folder, then move into place."""
    fd, name = tempfile.mkstemp(dir=out.parent, prefix=f".{out.name}.", suffix=".tmp")
    os.close(fd)
    tmp = Path(name)
    opts = dict(garbage=3, deflate=True)
    opts.update(kw)
    try:
        doc.save(tmp, **opts)
        os.replace(tmp, out)
    finally:
        if tmp.exists():
            tmp.unlink()
    return out


def _derotate(page, point):
    return pymupdf.Point(point) * page.derotation_matrix


# ---------------------------------------------------------------- organise
def organize(items: list[dict], output: str, title: str = "") -> dict:
    """Build a PDF from page references.

    item: {"path", "page", "rotate"} | {"image": path, "rotate"} | {"blank": [w, h]}
    Rotation is in degrees, added to the page's own rotation.
    """
    inputs = [i.get("path") or i.get("image") for i in items if i.get("path") or i.get("image")]
    out_path = _check_output(output, inputs)
    out = pymupdf.open()
    cache: dict[str, object] = {}
    try:
        for it in items:
            if "blank" in it:
                w, h = it["blank"]
                out.new_page(width=w, height=h)
                continue
            src_path = it.get("path") or it.get("image")
            src = cache.get(src_path)
            if src is None:
                src = _as_pdf(src_path)
                cache[src_path] = src
            pno = int(it.get("page", 0))
            if not 0 <= pno < src.page_count:
                raise ToolError(f"“{Path(src_path).name}” has no page {pno + 1}.")
            out.insert_pdf(src, from_page=pno, to_page=pno)
            rot = int(it.get("rotate", 0)) % 360
            if rot:
                page = out[-1]
                page.set_rotation((page.rotation + rot) % 360)
        if out.page_count == 0:
            raise ToolError("There are no pages to save.")
        if title:
            out.set_metadata({"title": title, "creator": "ENG Drawing Viewer"})
        _save(out, out_path, garbage=4)
    finally:
        for d in cache.values():
            d.close()
        out.close()
    return {"outputs": [str(out_path)], "summary": f"Saved {len(items)} page{'s' if len(items) != 1 else ''}."}


def _page_count(path: str) -> int:
    doc = _open(path)
    try:
        return doc.page_count
    finally:
        doc.close()


def merge(paths: list[str], output: str, bookmarks: bool = True) -> dict:
    out_path = _check_output(output, paths)
    out = pymupdf.open()
    toc = []
    try:
        for p in paths:
            src = _as_pdf(p)
            try:
                if bookmarks:
                    toc.append([1, Path(p).stem, out.page_count + 1])
                out.insert_pdf(src)
            finally:
                src.close()
        if bookmarks and toc:
            out.set_toc(toc)
        _save(out, out_path, garbage=4)
        n = out.page_count
    finally:
        out.close()
    return {"outputs": [str(out_path)], "summary": f"Merged {len(paths)} files into {n} pages."}


def split(path: str, out_dir: str, mode: str = "each", ranges: str = "", every: int = 1) -> dict:
    """mode: 'each' page, 'every' N pages, or 'ranges' ('1-3, 4-6' → one file per group)."""
    src = _as_pdf(path)
    stem = Path(path).stem
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    groups: list[list[int]] = []
    n = src.page_count
    if mode == "each":
        groups = [[i] for i in range(n)]
    elif mode == "every":
        k = max(1, int(every))
        groups = [list(range(i, min(i + k, n))) for i in range(0, n, k)]
    else:
        for part in [p for p in re.split(r"[;,]", ranges or "") if p.strip()]:
            groups.append(parse_ranges(part, n))
        if not groups:
            raise ToolError("Enter page ranges, for example 1-3, 4-8.")
    outputs = []
    try:
        for g in groups:
            label = f"p{g[0] + 1}" if len(g) == 1 else f"p{g[0] + 1}-{g[-1] + 1}"
            target = unique_path(folder / f"{stem} {label}.pdf")
            _check_output(str(target), [path])
            part = pymupdf.open()
            for i in g:
                part.insert_pdf(src, from_page=i, to_page=i)
            _save(part, target, garbage=4)
            part.close()
            outputs.append(str(target))
    finally:
        src.close()
    return {"outputs": outputs, "summary": f"Created {len(outputs)} PDF{'s' if len(outputs) != 1 else ''}."}


def select_pages(path: str, output: str, ranges: str, keep: bool = True) -> dict:
    """Extract (keep=True) or remove (keep=False) the given pages."""
    n = _page_count(path)
    chosen = parse_ranges(ranges, n)
    pages = chosen if keep else [i for i in range(n) if i not in set(chosen)]
    if not pages:
        raise ToolError("That would leave no pages.")
    res = organize([{"path": path, "page": i} for i in pages], output)
    res["summary"] = f"{'Extracted' if keep else 'Kept'} {len(pages)} of {n} pages."
    return res


def rotate(path: str, output: str, angle: int = 90, ranges: str = "") -> dict:
    n = _page_count(path)
    chosen = set(parse_ranges(ranges, n))
    items = [{"path": path, "page": i, "rotate": angle if i in chosen else 0} for i in range(n)]
    res = organize(items, output)
    res["summary"] = f"Rotated {len(chosen)} page{'s' if len(chosen) != 1 else ''} by {angle}°."
    return res


# ---------------------------------------------------------------- convert
def images_to_pdf(paths: list[str], output: str, page_size: str = "fit", landscape: str = "auto",
                  margin_mm: float = 0.0, separate: bool = False, out_dir: str | None = None) -> dict:
    """Images (and PDFs) into one PDF, or one PDF per image."""
    def place(dst, img_pdf):
        src_rect = img_pdf[0].rect
        if page_size == "fit":
            dst.insert_pdf(img_pdf)
            return
        land = src_rect.width > src_rect.height if landscape == "auto" else landscape == "landscape"
        w, h = paper_size(page_size, land)
        page = dst.new_page(width=w, height=h)
        m = margin_mm * MM
        page.show_pdf_page(pymupdf.Rect(m, m, w - m, h - m), img_pdf, 0, keep_proportion=True)

    outputs = []
    if separate:
        folder = Path(out_dir or Path(paths[0]).parent)
        for p in paths:
            doc = pymupdf.open()
            src = _as_pdf(p)
            for i in range(src.page_count):
                one = pymupdf.open()
                one.insert_pdf(src, from_page=i, to_page=i)
                place(doc, one)
                one.close()
            src.close()
            target = unique_path(folder / f"{Path(p).stem}.pdf")
            _check_output(str(target), paths)
            _save(doc, target, garbage=4)
            doc.close()
            outputs.append(str(target))
    else:
        out_path = _check_output(output, paths)
        doc = pymupdf.open()
        for p in paths:
            src = _as_pdf(p)
            for i in range(src.page_count):
                one = pymupdf.open()
                one.insert_pdf(src, from_page=i, to_page=i)
                place(doc, one)
                one.close()
            src.close()
        _save(doc, out_path, garbage=4)
        doc.close()
        outputs.append(str(out_path))
    return {"outputs": outputs, "summary": f"Converted {len(paths)} file{'s' if len(paths) != 1 else ''}."}


def to_images(path: str, out_dir: str, fmt: str = "png", dpi: int = 200, ranges: str = "") -> dict:
    doc = _open(path)
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    stem = Path(path).stem
    outputs = []
    try:
        pages = parse_ranges(ranges, doc.page_count)
        lowered = []
        for i in pages:
            page_dpi = max_dpi_for(doc[i], int(dpi))
            if page_dpi < int(dpi):
                lowered.append((i + 1, page_dpi))
            pix = render_page(doc[i], dpi=page_dpi, alpha=False)
            target = unique_path(folder / f"{stem} p{i + 1}.{fmt}")
            _check_output(str(target), [path])
            if fmt == "jpg":
                pix.set_dpi(page_dpi, page_dpi)
                pix.save(target, output="jpg", jpg_quality=92)
            elif fmt == "tiff":
                from PIL import Image

                Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(target, dpi=(page_dpi, page_dpi), compression="tiff_lzw")
            else:
                pix.set_dpi(page_dpi, page_dpi)
                pix.save(target, output="png")
            outputs.append(str(target))
    finally:
        doc.close()
    summary = f"Saved {len(outputs)} image{'s' if len(outputs) != 1 else ''} at {dpi} DPI."
    if lowered:
        detail = ", ".join(f"page {n} at {d} DPI" for n, d in lowered[:5]) + (" …" if len(lowered) > 5 else "")
        summary += f" Very large pages were saved at a lower resolution to stay within memory limits: {detail}."
    return {"outputs": outputs, "summary": summary}


def to_text(path: str, output: str, markdown: bool = False) -> dict:
    out_path = _check_output(output, [path])
    doc = _open(path)
    parts = []
    try:
        for i, page in enumerate(doc):
            text = page.get_text("text", sort=True).strip()
            if markdown:
                parts.append(f"## Page {i + 1}\n\n{text}\n")
            else:
                parts.append(text + "\n\f" if i < doc.page_count - 1 else text + "\n")
        title = Path(path).stem
    finally:
        doc.close()
    body = (f"# {title}\n\n" if markdown else "") + "\n".join(parts)
    out_path.write_text(body, encoding="utf-8")
    words = len(body.split())
    return {"outputs": [str(out_path)], "summary": f"Extracted about {words} words." if words else
            "No text found — this looks like a scan. Run OCR first to make it searchable."}


def to_word(path: str, output: str, mode: str = "text") -> dict:
    """mode: 'text' (editable paragraphs), 'images' (one picture per page) or 'both'."""
    from docx import Document
    from docx.shared import Pt, Mm

    out_path = _check_output(output, [path])
    src = _open(path)
    docx = Document()
    sec = docx.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Mm(15)
    usable_w = sec.page_width - sec.left_margin - sec.right_margin
    try:
        for i, page in enumerate(src):
            if i:
                docx.add_page_break()
            if mode in ("images", "both"):
                pix = render_page(page, dpi=150, alpha=False)
                docx.add_picture(io.BytesIO(pix.tobytes("png")), width=usable_w)
            if mode in ("text", "both"):
                data = page.get_text("dict", sort=True)
                for block in data["blocks"]:
                    if block.get("type") != 0:
                        continue
                    para = docx.add_paragraph()
                    for line in block["lines"]:
                        for span in line["spans"]:
                            t = span["text"]
                            if not t.strip():
                                continue
                            run = para.add_run(t + " ")
                            run.font.size = Pt(max(6, min(36, round(span["size"]))))
                            run.bold = bool(span["flags"] & 16)
                            run.italic = bool(span["flags"] & 2)
    finally:
        src.close()
    with _temporary_output(out_path) as tmp:
        docx.save(tmp)
    return {"outputs": [str(out_path)], "summary": f"Created a Word document from {i + 1} page{'s' if i else ''}."}


def to_excel(path: str, output: str, include_text: bool = True) -> dict:
    """Tables (e.g. BOMs and title blocks) to sheets; optionally all text lines."""
    import xlsxwriter

    out_path = _check_output(output, [path])
    with _temporary_output(out_path, ".xlsx") as tmp:
        wb = xlsxwriter.Workbook(str(tmp), {"strings_to_formulas": False, "strings_to_urls": False})
        bold = wb.add_format({"bold": True, "bg_color": "#EEF2F7", "border": 1})
        cell = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
        src = _open(path)
        n_tables = 0
        try:
            for i, page in enumerate(src):
                try:
                    tables = page.find_tables().tables
                except Exception:
                    tables = []
                for k, table in enumerate(tables):
                    rows = table.extract()
                    if not rows:
                        continue
                    ws = wb.add_worksheet(f"P{i + 1} T{k + 1}"[:31])
                    for r, row in enumerate(rows):
                        for c, val in enumerate(row):
                            ws.write(r, c, "" if val is None else str(val), bold if r == 0 else cell)
                    widths = [max((len(str(row[c] or "")) for row in rows if c < len(row)), default=8) for c in range(max(len(r) for r in rows))]
                    for c, w in enumerate(widths):
                        ws.set_column(c, c, min(60, max(8, w + 2)))
                    n_tables += 1
            if include_text or n_tables == 0:
                ws = wb.add_worksheet("Text")
                ws.write_row(0, 0, ["Page", "Text"], bold)
                r = 1
                for i, page in enumerate(src):
                    for line in page.get_text("text", sort=True).splitlines():
                        if line.strip():
                            ws.write_row(r, 0, [i + 1, line.strip()])
                            r += 1
                ws.set_column(1, 1, 90)
        finally:
            src.close()
            wb.close()
    return {"outputs": [str(out_path)], "summary": f"Found {n_tables} table{'s' if n_tables != 1 else ''}."
            + ("" if n_tables else " All text lines were exported instead.")}


def to_powerpoint(path: str, output: str, dpi: int = 200, title_mode: str = "title") -> dict:
    from . import pptx_export

    out_path = _check_output(output, [path])
    src = _open(path)
    tmpdir = Path(tempfile.mkdtemp(prefix="eng-pdf2pptx-"))
    try:
        slides = []
        for i, page in enumerate(src):
            pymupdf.TOOLS.set_graphics_min_line_width(max(1.0, 0.9 * dpi / 72))
            try:
                pix = render_page(page, dpi=dpi, alpha=False)
            finally:
                pymupdf.TOOLS.set_graphics_min_line_width(0)
            pix.set_dpi(dpi, dpi)
            png = tmpdir / f"p{i + 1}.png"
            pix.save(png)
            title = Path(path).name if src.page_count == 1 else f"{Path(path).name} — page {i + 1}"
            slides.append({"image": str(png), "img_w": pix.width, "img_h": pix.height, "title": title})
        info = pptx_export.build({"output": str(out_path), "mode": "new", "title_mode": title_mode, "slides": slides})
    finally:
        src.close()
        shutil.rmtree(tmpdir, ignore_errors=True)
    return {"outputs": [str(out_path)], "summary": f"Created {info['slides']} slides."}


def extract_images(path: str, out_dir: str, min_px: int = 32) -> dict:
    doc = _open(path)
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    seen = set()
    outputs = []
    try:
        if not doc.is_pdf:
            raise ToolError("This file is already an image.")
        for i, page in enumerate(doc):
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen:
                    continue
                seen.add(xref)
                check_pixels(img[2], img[3])
                info = doc.extract_image(xref)
                if not info or info.get("width", 0) < min_px or info.get("height", 0) < min_px:
                    continue
                target = unique_path(folder / f"{Path(path).stem} p{i + 1} image {len(outputs) + 1}.{info['ext']}")
                target.write_bytes(info["image"])
                outputs.append(str(target))
    finally:
        doc.close()
    return {"outputs": outputs, "summary": f"Extracted {len(outputs)} image{'s' if len(outputs) != 1 else ''}."
            if outputs else "This PDF contains no embedded images (drawings are usually vector)."}


def office_to_pdf(paths: list[str], out_dir: str) -> dict:
    """Word/Excel/PowerPoint/OpenDocument → PDF with a local LibreOffice (offline)."""
    soffice = find_office()
    if not soffice:
        raise ToolError("Converting Office documents needs LibreOffice (free) installed on this computer.")
    folder = Path(out_dir)
    folder.mkdir(parents=True, exist_ok=True)
    outputs = []
    with tempfile.TemporaryDirectory(prefix="eng-office-") as tmp:
        profile = Path(tmp, "profile").as_uri()
        for p in paths:
            work = Path(tmp, f"out-{len(outputs)}")
            work.mkdir(exist_ok=True)
            command = [str(Path(soffice).resolve()), f"-env:UserInstallation={profile}",
                       "--headless", "--norestore", "--nodefault", "--nolockcheck",
                       "--convert-to", "pdf", "--outdir", str(work),
                       str(Path(p).expanduser().resolve())]
            try:
                returncode = run_conversion(command, timeout=300)
            except subprocess.TimeoutExpired:
                raise ToolError("LibreOffice conversion timed out; its contained processes were stopped.") from None
            except OSError as exc:
                raise ToolError(f"Could not run LibreOffice: {exc}") from None
            produced = work / (Path(p).stem + ".pdf")
            if returncode != 0 or not produced.is_file() or produced.is_symlink():
                raise ToolError(f"LibreOffice could not convert “{Path(p).name}”.")
            target = unique_path(folder / produced.name)
            _check_output(str(target), paths)
            with _temporary_output(target) as temporary:
                shutil.copyfile(produced, temporary)
            outputs.append(str(target))
    return {"outputs": outputs, "summary": f"Converted {len(outputs)} document{'s' if len(outputs) != 1 else ''}."}


# ---------------------------------------------------------------- optimise
COMPRESSION = {
    "lossless": None,
    "balanced": dict(dpi_threshold=220, dpi_target=150, quality=80),
    "strong": dict(dpi_threshold=130, dpi_target=96, quality=60),
}


def compress(path: str, output: str, level: str = "balanced") -> dict:
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    before = os.path.getsize(path)
    try:
        opts = COMPRESSION.get(level)
        if opts:
            try:
                doc.rewrite_images(**opts)
            except Exception:
                pass  # images that cannot be recompressed are kept as they are
        try:
            doc.subset_fonts()
        except Exception:
            pass
        _save(doc, out_path, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True, use_objstms=1)
    finally:
        doc.close()
    after = os.path.getsize(out_path)
    pct = (1 - after / before) * 100 if before else 0
    note = f"{_fmt_size(before)} → {_fmt_size(after)}" + (f" ({pct:.0f}% smaller)" if pct > 0.5 else " (already compact)")
    return {"outputs": [str(out_path)], "summary": note, "before": before, "after": after}


def _fmt_size(n: int) -> str:
    for unit in ("bytes", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "bytes" else f"{n:.1f} {unit}"
        n /= 1024
    return str(n)


def repair(path: str, output: str) -> dict:
    out_path = _check_output(output, [path])
    try:
        doc = pymupdf.open(path, filetype="pdf")  # MuPDF rebuilds broken cross-reference tables while opening
    except Exception:
        raise ToolError("This file is too damaged to recover.") from None
    try:
        if doc.needs_pass:
            raise ToolError("Unlock the PDF before repairing it.")
        n = doc.page_count
        if n == 0:
            raise ToolError("No pages could be recovered.")
        _save(doc, out_path, garbage=4, deflate=True, clean=True)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": f"Recovered {n} page{'s' if n != 1 else ''}."}


def flatten(path: str, output: str) -> dict:
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    try:
        doc.bake(annots=True, widgets=True)
        _save(doc, out_path, garbage=4)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": "Comments, markups and form fields are now part of the page content."}


# ---- OCR ---------------------------------------------------------------
def ocr_engine() -> str | None:
    if sys.platform == "darwin":
        try:
            import Vision  # noqa: F401

            return "vision"
        except Exception:
            pass
    if shutil.which("tesseract"):
        return "tesseract"
    return None


def _vision_lines(png: bytes, languages: list[str] | None, fast: bool):
    import Vision
    from Foundation import NSData

    handler = Vision.VNImageRequestHandler.alloc().initWithData_options_(NSData.dataWithBytes_length_(png, len(png)), None)
    req = Vision.VNRecognizeTextRequest.alloc().init()
    req.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelFast if fast else Vision.VNRequestTextRecognitionLevelAccurate)
    req.setUsesLanguageCorrection_(True)
    if languages:
        req.setRecognitionLanguages_(languages)
    ok, err = handler.performRequests_error_([req], None)
    if not ok:
        raise ToolError(f"Text recognition failed: {err}")
    out = []
    for obs in req.results() or []:
        cand = obs.topCandidates_(1)
        if not cand:
            continue
        text = str(cand[0].string())
        bb = obs.boundingBox()  # normalised, origin bottom-left
        out.append((text, bb.origin.x, bb.origin.y, bb.size.width, bb.size.height))
    return out


def _tesseract_lines(page, dpi: int, language: str):
    tp = page.get_textpage_ocr(dpi=dpi, full=True, language=language or "eng")
    out = []
    r = page.rect
    for x0, y0, x1, y1, word, *_ in page.get_text("words", textpage=tp):
        out.append((word, x0 / r.width, 1 - y1 / r.height, (x1 - x0) / r.width, (y1 - y0) / r.height))
    return out


def ocr(path: str, output: str, languages: str = "", skip_text_pages: bool = True, dpi: int = 300) -> dict:
    """Add an invisible, searchable text layer. Pages keep their look exactly."""
    engine = ocr_engine()
    if engine is None:
        raise ToolError("Text recognition needs macOS, or Tesseract OCR installed on this computer.")
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    langs = [l.strip() for l in languages.split(",") if l.strip()] or None
    font = pymupdf.Font("helv")
    added = pages_done = 0
    try:
        for page in doc:
            if skip_text_pages and page.get_text("text").strip():
                continue
            if engine == "vision":
                png = render_page(page, dpi=dpi, alpha=False).tobytes("png")
                lines = _vision_lines(png, langs, fast=False)
            else:
                lines = _tesseract_lines(page, dpi, "+".join(langs) if langs else "eng")
            r = page.rect
            for text, nx, ny, nw, nh in lines:
                h = nh * r.height
                if h <= 0 or not text.strip():
                    continue
                fs = h * 0.82
                x0 = nx * r.width
                baseline = (1 - ny) * r.height - h * 0.18
                width = font.text_length(text, fontsize=fs)
                sx = (nw * r.width) / width if width > 0 else 1.0
                p = _derotate(page, (x0, baseline))
                page.insert_text(p, text, fontsize=fs, fontname="helv", render_mode=3, rotate=page.rotation,
                                 morph=(p, pymupdf.Matrix(sx, 1)))
                added += 1
            pages_done += 1
        _save(doc, out_path, garbage=4)
    finally:
        doc.close()
    if pages_done == 0:
        return {"outputs": [str(out_path)], "summary": "Every page already had text — nothing needed recognising."}
    return {"outputs": [str(out_path)], "summary": f"Recognised {added} lines of text on {pages_done} page{'s' if pages_done != 1 else ''}; the PDF is now searchable."}


# ---------------------------------------------------------------- edit
POSITIONS = ("bottom-center", "bottom-right", "bottom-left", "top-center", "top-right", "top-left")


def _anchor(rect, position: str, margin: float, tw: float, fs: float) -> tuple[float, float]:
    v, h = position.split("-")
    x = {"left": rect.x0 + margin, "center": rect.x0 + (rect.width - tw) / 2, "right": rect.x1 - margin - tw}[h]
    y = rect.y0 + margin + fs if v == "top" else rect.y1 - margin
    return x, y


def page_numbers(path: str, output: str, position: str = "bottom-right", fmt: str = "Page {n} of {total}",
                 start: int = 1, size: float = 10, margin_mm: float = 8, ranges: str = "", color: str = "#000000") -> dict:
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    rgb = _rgb(color)
    font = pymupdf.Font("helv")
    try:
        chosen = parse_ranges(ranges, doc.page_count)
        for k, i in enumerate(chosen):
            page = doc[i]
            text = fmt.replace("{n}", str(start + k)).replace("{total}", str(len(chosen) + start - 1)).replace("{file}", Path(path).stem)
            tw = font.text_length(text, fontsize=size)
            x, y = _anchor(page.rect, position, margin_mm * MM, tw, size)
            page.insert_text(_derotate(page, (x, y)), text, fontsize=size, fontname="helv", color=rgb, rotate=page.rotation)
        _save(doc, out_path)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": f"Numbered {len(chosen)} pages."}


def _rgb(hexcolor: str) -> tuple[float, float, float]:
    h = hexcolor.lstrip("#")
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)) if len(h) == 6 else (0, 0, 0)


STAMPS = ["DRAFT", "PRELIMINARY", "FOR REVIEW", "APPROVED", "REJECTED", "SUPERSEDED", "VOID", "CONFIDENTIAL", "FOR INFORMATION ONLY", "AS BUILT"]


def watermark(path: str, output: str, text: str = "DRAFT", style: str = "diagonal", color: str = "#D21E1E",
              opacity: float = 0.25, size: float = 0, add_date: bool = False, behind: bool = False,
              image: str = "", ranges: str = "") -> dict:
    """style: 'diagonal' watermark across the page, or 'stamp' (boxed, top right).

    With ``image`` set, the image is placed (centred, scaled) instead of text.
    """
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    rgb = _rgb(color)
    label = text.strip()
    if add_date:
        label = f"{label}  {_dt.date.today().isoformat()}"
    font = pymupdf.Font("hebo")
    img_bytes = None
    if image:
        img_bytes = _image_with_opacity(image, opacity)
    try:
        chosen = parse_ranges(ranges, doc.page_count)
        for i in chosen:
            page = doc[i]
            r = page.rect
            if img_bytes is not None:
                w = r.width * 0.5
                rect = pymupdf.Rect(r.x0 + (r.width - w) / 2, r.y0 + r.height * 0.25, r.x0 + (r.width + w) / 2, r.y1 - r.height * 0.25)
                page.insert_image(rect * page.derotation_matrix, stream=img_bytes, keep_proportion=True,
                                  overlay=not behind, rotate=page.rotation)
                continue
            if style == "stamp":
                fs = size or max(10.0, min(r.width, r.height) * 0.025)
                tw = font.text_length(label, fontsize=fs)
                pad = fs * 0.5
                box = pymupdf.Rect(r.x1 - tw - 2 * pad - 20, r.y0 + 20, r.x1 - 20, r.y0 + 20 + fs + 2 * pad)
                shape = page.new_shape()
                shape.draw_rect(box * page.derotation_matrix)
                shape.finish(color=rgb, width=max(1.0, fs / 10), stroke_opacity=min(1, opacity * 3))
                shape.commit(overlay=not behind)
                p = _derotate(page, (box.x0 + pad, box.y1 - pad - fs * 0.15))
                page.insert_text(p, label, fontsize=fs, fontname="hebo", fontfile=None, color=rgb,
                                 fill_opacity=min(1, opacity * 3), rotate=page.rotation, overlay=not behind)
            else:
                diag = (r.width ** 2 + r.height ** 2) ** 0.5
                fs = size or max(24.0, diag / max(len(label), 4) * 0.95)
                tw = font.text_length(label, fontsize=fs)
                import math

                ang = math.degrees(math.atan2(r.height, r.width))
                cx, cy = r.x0 + r.width / 2, r.y0 + r.height / 2
                start = pymupdf.Point(cx - tw / 2, cy + fs * 0.35)
                centre = pymupdf.Point(cx, cy)
                p = _derotate(page, start)
                pivot = _derotate(page, centre)
                page.insert_text(p, label, fontsize=fs, fontname="hebo", color=rgb, fill_opacity=opacity,
                                 rotate=page.rotation, overlay=not behind, morph=(pivot, pymupdf.Matrix(ang)))
        _save(doc, out_path)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": f"Stamped {len(chosen)} page{'s' if len(chosen) != 1 else ''}."}


def _image_with_opacity(path: str, opacity: float) -> bytes:
    from PIL import Image

    if extension(path) in (".heic", ".heif"):
        from .documents import heic_to_png

        im = Image.open(io.BytesIO(heic_to_png(path)))
    else:
        im = Image.open(path)
    check_pixels(*im.size)
    im = im.convert("RGBA")
    if opacity < 1:
        alpha = im.getchannel("A").point(lambda a: int(a * max(0.0, min(1.0, opacity))))
        im.putalpha(alpha)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def _content_bbox(page) -> pymupdf.Rect | None:
    r = page.rect
    box = pymupdf.Rect()
    for kind, rect in page.get_bboxlog():
        rect = pymupdf.Rect(rect)
        if rect.is_empty or rect.is_infinite:
            continue
        if kind == "fill-path" and rect.width >= r.width * 0.98 and rect.height >= r.height * 0.98:
            continue  # a page-sized background fill is not content
        if kind.startswith(("fill-", "stroke-")):
            box |= rect
    return None if box.is_empty else box & r


def crop(path: str, output: str, mode: str = "auto", margin_mm: float = 5.0,
         top: float = 0, right: float = 0, bottom: float = 0, left: float = 0) -> dict:
    """mode 'auto': trim to the drawn content plus a margin; 'margins': cut fixed mm from each edge."""
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    changed = 0
    try:
        for page in doc:
            r = page.rect
            if mode == "auto":
                bbox = _content_bbox(page)
                if bbox is None:
                    continue
                m = margin_mm * MM
                target = pymupdf.Rect(bbox.x0 - m, bbox.y0 - m, bbox.x1 + m, bbox.y1 + m) & r
            else:
                target = pymupdf.Rect(r.x0 + left * MM, r.y0 + top * MM, r.x1 - right * MM, r.y1 - bottom * MM)
            if target.is_empty or target.width < 10 or target.height < 10:
                raise ToolError("Those margins would leave nothing of the page.")
            # set_cropbox works in unrotated page space, relative to the mediabox.
            unrot = target * page.derotation_matrix
            mb = page.mediabox
            page.set_cropbox(pymupdf.Rect(unrot.x0 + mb.x0, unrot.y0 + mb.y0, unrot.x1 + mb.x0, unrot.y1 + mb.y0) & mb)
            changed += 1
        _save(doc, out_path)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": f"Cropped {changed} page{'s' if changed != 1 else ''}."}


def resize(path: str, output: str, paper: str = "A3", orientation: str = "auto", margin_mm: float = 0.0) -> dict:
    """Scale every page onto a standard sheet (vector content stays vector)."""
    out_path = _check_output(output, [path])
    src = _as_pdf(path)
    out = pymupdf.open()
    try:
        for i, page in enumerate(src):
            r = page.rect
            land = r.width > r.height if orientation == "auto" else orientation == "landscape"
            w, h = paper_size(paper, land)
            dst = out.new_page(width=w, height=h)
            m = margin_mm * MM
            dst.show_pdf_page(pymupdf.Rect(m, m, w - m, h - m), src, i, keep_proportion=True)
        _save(out, out_path, garbage=4)
        n = out.page_count
    finally:
        src.close()
        out.close()
    return {"outputs": [str(out_path)], "summary": f"Resized {n} page{'s' if n != 1 else ''} to {paper}."}


def tile(path: str, output: str, paper: str = "A4", orientation: str = "auto", scale_pct: float = 100.0,
         overlap_mm: float = 10.0, margin_mm: float = 10.0, labels: bool = True, ranges: str = "") -> dict:
    """Print a large sheet on small paper: split each page into overlapping tiles."""
    out_path = _check_output(output, [path])
    src = _as_pdf(path)
    out = pymupdf.open()
    sheets = 0
    try:
        for i in parse_ranges(ranges, src.page_count):
            r = src[i].rect
            land = r.width > r.height if orientation == "auto" else orientation == "landscape"
            pw, ph = paper_size(paper, land)
            m, ov = margin_mm * MM, overlap_mm * MM
            s = scale_pct / 100.0
            area_w, area_h = (pw - 2 * m) / s, (ph - 2 * m) / s  # source points per tile
            step_w, step_h = max(area_w - ov / s, 1), max(area_h - ov / s, 1)
            cols = max(1, int((r.width - ov / s) // step_w + (1 if (r.width - ov / s) % step_w > 1e-6 else 0)))
            rows = max(1, int((r.height - ov / s) // step_h + (1 if (r.height - ov / s) % step_h > 1e-6 else 0)))
            for row in range(rows):
                for col in range(cols):
                    clip = pymupdf.Rect(col * step_w, row * step_h, col * step_w + area_w, row * step_h + area_h) & r
                    dst = out.new_page(width=pw, height=ph)
                    target = pymupdf.Rect(m, m, m + clip.width * s, m + clip.height * s)
                    dst.show_pdf_page(target, src, i, clip=clip)
                    if labels:
                        shape = dst.new_shape()
                        for x, y in ((m, m), (m + clip.width * s, m), (m, m + clip.height * s), (m + clip.width * s, m + clip.height * s)):
                            shape.draw_line((x - 6, y), (x + 6, y))
                            shape.draw_line((x, y - 6), (x, y + 6))
                        shape.finish(color=(0.4, 0.4, 0.4), width=0.4)
                        shape.commit()
                        tag = f"{Path(path).stem} · p{i + 1} · row {row + 1}/{rows} col {col + 1}/{cols} · overlap {overlap_mm:g} mm"
                        dst.insert_text((m, ph - m / 2), tag, fontsize=6, color=(0.4, 0.4, 0.4))
                    sheets += 1
        _save(out, out_path, garbage=4)
    finally:
        src.close()
        out.close()
    return {"outputs": [str(out_path)], "summary": f"Split into {sheets} {paper} sheet{'s' if sheets != 1 else ''}."}


def nup(path: str, output: str, per_sheet: int = 4, paper: str = "A4", orientation: str = "auto", gap_mm: float = 5.0) -> dict:
    """Several pages on one sheet (handouts, overviews)."""
    grids = {2: (2, 1), 4: (2, 2), 6: (3, 2), 8: (4, 2), 9: (3, 3), 16: (4, 4)}
    if per_sheet not in grids:
        raise ToolError("Choose 2, 4, 6, 8, 9 or 16 pages per sheet.")
    out_path = _check_output(output, [path])
    src = _as_pdf(path)
    out = pymupdf.open()
    try:
        first = src[0].rect
        land = first.width > first.height if orientation == "auto" else orientation == "landscape"
        cols, rows = grids[per_sheet]
        if not land:
            cols, rows = rows, cols
        pw, ph = paper_size(paper, land)
        g = gap_mm * MM
        cw, ch = (pw - g * (cols + 1)) / cols, (ph - g * (rows + 1)) / rows
        page = None
        for i in range(src.page_count):
            k = i % per_sheet
            if k == 0:
                page = out.new_page(width=pw, height=ph)
            c, r = k % cols, k // cols
            cell = pymupdf.Rect(g + c * (cw + g), g + r * (ch + g), g + c * (cw + g) + cw, g + r * (ch + g) + ch)
            page.show_pdf_page(cell, src, i, keep_proportion=True)
        _save(out, out_path, garbage=4)
        n = out.page_count
    finally:
        src.close()
        out.close()
    return {"outputs": [str(out_path)], "summary": f"Placed {per_sheet} pages per sheet on {n} sheet{'s' if n != 1 else ''}."}


def overlay(path: str, overlay_path: str, output: str, mode: str = "first", behind: bool = False) -> dict:
    """Stamp another PDF (e.g. a title block or approval stamp) onto every page.

    mode 'first': the overlay's first page on every page; 'match': page n on page n.
    """
    out_path = _check_output(output, [path, overlay_path])
    doc = _as_pdf(path)
    ov = _as_pdf(overlay_path)
    try:
        for i, page in enumerate(doc):
            j = 0 if mode == "first" else min(i, ov.page_count - 1)
            page.show_pdf_page(page.rect * page.derotation_matrix, ov, j, overlay=not behind, rotate=-page.rotation)
        _save(doc, out_path, garbage=4)
        n = doc.page_count
    finally:
        doc.close()
        ov.close()
    return {"outputs": [str(out_path)], "summary": f"Overlaid {n} page{'s' if n != 1 else ''}."}


def _remove_auxiliary_data(doc) -> None:
    """Remove defined non-page-content channels; deliberately not a sanitizer.

    This is destructive: bookmarks, comments, attachments, links and form fields
    are removed in full, rather than selectively matching a redaction term.
    Page content (including images, OCR layers and arbitrary PDF constructs)
    still requires visual and independent inspection before sharing.
    """
    doc.set_toc([])
    for page in doc:
        # Capture identifiers first: deleting a parent can also delete its popup.
        for xref in [annot.xref for annot in (page.annots() or ())]:
            annot = page.load_annot(xref)
            if annot is not None:
                page.delete_annot(annot)
        for widget in list(page.widgets() or ()):
            page.delete_widget(widget)
    doc.scrub(metadata=True, xml_metadata=True, attached_files=True,
              embedded_files=True, clean_pages=True, hidden_text=False,
              javascript=True, redactions=False, remove_links=True,
              reset_fields=False, reset_responses=False, thumbnails=True)
    catalog = doc.pdf_catalog()
    # Remove document-level containers that can carry names, form defaults,
    # accessibility text, actions or application-private information.
    for key in ("Names", "AcroForm", "Outlines", "OpenAction", "AA",
                "PieceInfo", "Collection", "StructTreeRoot", "MarkInfo", "PageLabels"):
        doc.xref_set_key(catalog, key, "null")
    for page in doc:
        for key in ("AA", "PieceInfo", "PresSteps"):
            doc.xref_set_key(page.xref, key, "null")


def metadata(path: str, output: str, title: str = "", author: str = "", subject: str = "", keywords: str = "",
             strip: bool = False) -> dict:
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    try:
        if strip:
            _remove_auxiliary_data(doc)
        else:
            meta = dict(doc.metadata or {})
            meta.update({k: v for k, v in (("title", title), ("author", author), ("subject", subject), ("keywords", keywords)) if v})
            doc.set_metadata(meta)
        _save(doc, out_path, garbage=4)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": "Removed document metadata, bookmarks, comments, attachments, links and form fields. Page content still needs review." if strip else "Metadata updated."}


# ---------------------------------------------------------------- security
def protect(path: str, output: str, user_password: str, owner_password: str = "", allow_print: bool = True,
            allow_copy: bool = False, allow_edit: bool = False) -> dict:
    if not user_password:
        raise ToolError("Enter a password.")
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    perm = pymupdf.PDF_PERM_ACCESSIBILITY
    if allow_print:
        perm |= pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ
    if allow_copy:
        perm |= pymupdf.PDF_PERM_COPY
    if allow_edit:
        perm |= pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM | pymupdf.PDF_PERM_ASSEMBLE
    try:
        _save(doc, out_path, garbage=4, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_password,
              owner_pw=owner_password or secrets.token_hex(20), permissions=perm)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": "Encrypted with AES-256. Keep the password safe — it cannot be recovered."}


def unlock(path: str, output: str, password: str) -> dict:
    out_path = _check_output(output, [path])
    doc = _open(path, password=password)
    try:
        _save(doc, out_path, garbage=4, encryption=pymupdf.PDF_ENCRYPT_NONE)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": "Saved an unlocked copy."}


def redact(path: str, output: str, terms: str, match_case: bool = False, whole_words: bool = False,
           fill: str = "#000000") -> dict:
    """Remove matching searchable page text; not a document sanitization guarantee."""
    words = [t.strip() for t in re.split(r"[\n,;]", terms or "") if t.strip()]
    if not words:
        raise ToolError("Enter the words or numbers to redact.")
    out_path = _check_output(output, [path])
    doc = _as_pdf(path)
    count = 0
    try:
        for page in doc:
            hits = []
            for w in words:
                for rect in page.search_for(w):
                    if whole_words:
                        around = page.get_textbox(rect + (-3, 0, 3, 0))
                        if not re.search(rf"(?<![\w]){re.escape(w)}(?![\w])", around, 0 if match_case else re.I):
                            continue
                    if match_case and w not in page.get_textbox(rect):
                        continue
                    hits.append(rect)
            for rect in hits:
                page.add_redact_annot(rect, fill=_rgb(fill))
            if hits:
                page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS, graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED)
                count += len(hits)
        _remove_auxiliary_data(doc)
        _save(doc, out_path, garbage=4, clean=True)
    finally:
        doc.close()
    return {"outputs": [str(out_path)], "summary": f"Removed {count} page-text occurrence{'s' if count != 1 else ''}. Removed metadata, bookmarks, comments, attachments, links and form fields. Review images and remaining page content." if count
            else "No searchable text matched. Removed metadata, bookmarks, comments, attachments, links and form fields; review page content."}


# ---------------------------------------------------------------- compare
def text_changes(path_a: str, path_b: str, page_a: int = 0, page_b: int = 0, limit: int = 200) -> list[tuple[str, str]]:
    """Line-level differences between the text of two pages: [('-'|'+', line)]."""
    da, db = _open(path_a), _open(path_b)
    try:
        ta = [l.strip() for l in da[page_a].get_text("text", sort=True).splitlines() if l.strip()]
        tb = [l.strip() for l in db[page_b].get_text("text", sort=True).splitlines() if l.strip()]
    finally:
        da.close()
        db.close()
    out = []
    for line in difflib.unified_diff(ta, tb, lineterm="", n=0):
        if line.startswith(("---", "+++", "@@")):
            continue
        if line[:1] in "+-":
            out.append((line[0], line[1:]))
            if len(out) >= limit:
                break
    return out


def mark_pdf(source: str, page: int, overlay_pdf: str, output: str) -> dict:
    """One page plus a vector markup overlay (from MARK) as a new PDF."""
    out_path = _check_output(output, [source])
    src = _as_pdf(source)
    ov = pymupdf.open(overlay_pdf)
    out = pymupdf.open()
    try:
        out.insert_pdf(src, from_page=page, to_page=page)
        dst = out[0]
        dst.show_pdf_page(dst.rect * dst.derotation_matrix, ov, 0, overlay=True, rotate=dst.rotation)
        _save(out, out_path, garbage=4)
    finally:
        src.close()
        ov.close()
        out.close()
    return {"outputs": [str(out_path)], "summary": "Saved the marked page as a PDF. The original is unchanged."}


# ---------------------------------------------------------------- dispatch
OPS = {
    "organize": organize, "merge": merge, "split": split, "select_pages": select_pages, "rotate": rotate,
    "images_to_pdf": images_to_pdf, "to_images": to_images, "to_text": to_text, "to_word": to_word,
    "to_excel": to_excel, "to_powerpoint": to_powerpoint, "extract_images": extract_images,
    "office_to_pdf": office_to_pdf, "compress": compress, "repair": repair, "flatten": flatten, "ocr": ocr,
    "page_numbers": page_numbers, "watermark": watermark, "crop": crop, "resize": resize, "tile": tile, "nup": nup,
    "overlay": overlay, "metadata": metadata, "protect": protect, "unlock": unlock, "redact": redact,
    "mark_pdf": mark_pdf,
    "text_changes": lambda **kw: {"outputs": [], "changes": text_changes(**kw)},
}


def run(args: dict) -> dict:
    args = dict(args)
    op = args.pop("op")
    fn = OPS.get(op)
    if fn is None:
        raise ToolError(f"Unknown tool “{op}”.")
    try:
        return fn(**args)
    except ToolError:
        raise
    except TypeError as exc:
        raise ToolError(f"Internal option error: {exc}") from None
