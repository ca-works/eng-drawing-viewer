"""Build PowerPoint decks locally with python-pptx (no PowerPoint required).

Runs inside a render worker process. Layout follows the original COM export:
plain pages are centred; MARKed pages put the drawing on the right and the
numbered comments in an editable text box on the left. The PDF filename can be
an editable slide title or subtitle.
"""
from __future__ import annotations

import math
import os
import tempfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Inches, Pt

SLIDE_SIZES = {"16:9": (Inches(13.333), Inches(7.5)), "4:3": (Inches(10), Inches(7.5))}


def _layout(prs, want_title: bool):
    """Pick a 'Title Only' layout when a title is wanted, else the emptiest one."""
    best, best_score = None, None
    for layout in prs.slide_layouts:
        types = [ph.placeholder_format.type for ph in layout.placeholders]
        has_title = any(t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) for t in types)
        others = [t for t in types if t not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE,
                                                 PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)]
        if want_title:
            score = (0 if has_title else 1, len(others))
        else:
            score = (1 if has_title else 0, len(others))
        if best_score is None or score < best_score:
            best, best_score = layout, score
    return best


def _remove_empty_placeholders(slide, keep_title: bool):
    for ph in list(slide.placeholders):
        t = ph.placeholder_format.type
        if keep_title and t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            continue
        if ph.has_text_frame and ph.text_frame.text.strip():
            continue
        if t in (PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE,
                 PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            ph._element.getparent().remove(ph._element)


def _place_filename(slide, text: str, mode: str, sw: int, sh: int, new_slide: bool):
    if not text or mode == "none":
        return
    if mode == "title":
        title = slide.shapes.title
        if title is not None:
            title.text_frame.text = text
            if new_slide:
                title.left, title.top, title.width, title.height = int(sw * 0.05), int(sh * 0.03), int(sw * 0.90), int(sh * 0.10)
                for p in title.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(26)
            return
        box = slide.shapes.add_textbox(int(sw * 0.05), int(sh * 0.035), int(sw * 0.90), int(sh * 0.08))
        box.text_frame.text = text
        run = box.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(24)
        run.font.bold = True
        return
    for ph in slide.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE:
            ph.text_frame.text = text
            return
    box = slide.shapes.add_textbox(int(sw * 0.05), int(sh * 0.055), int(sw * 0.90), int(sh * 0.055))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(18)


def _comment_font_size(comments: list[str], box_w_pt: float, box_h_pt: float) -> int:
    for size in (18, 16, 14, 13, 12, 11, 10, 9):
        chars_per_line = max(8, int(box_w_pt / (size * 0.5)))
        lines = sum(max(1, math.ceil(len(c) / chars_per_line)) + c.count("\n") for c in comments) or 1
        if lines * size * 1.25 <= box_h_pt:
            return size
    return 9


def _add_comments(slide, comments: list[str], left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    size = _comment_font_size(comments, Emu(width).pt, Emu(height).pt)
    first = True
    for line in comments:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        num, sep, rest = line.partition(". ")
        if sep and num.isdigit():
            r1 = para.add_run()
            r1.text = f"{num}. "
            r1.font.bold = True
            r1.font.size = Pt(size)
            r2 = para.add_run()
            r2.text = rest
            r2.font.size = Pt(size)
        else:
            r = para.add_run()
            r.text = line
            r.font.size = Pt(size)
        para.space_after = Pt(size * 0.35)
    return box


def _fit(iw, ih, max_w, max_h):
    if iw <= 0 or ih <= 0:
        return max_w, max_h
    s = min(max_w / iw, max_h / ih)
    return int(iw * s), int(ih * s)


def build(spec: dict) -> dict:
    mode = spec.get("mode", "new")
    if mode == "existing":
        prs = Presentation(spec["template"])
    else:
        prs = Presentation()
        w, h = SLIDE_SIZES.get(spec.get("slide_size", "16:9"), SLIDE_SIZES["16:9"])
        prs.slide_width, prs.slide_height = w, h
    sw, sh = prs.slide_width, prs.slide_height
    title_mode = spec.get("title_mode", "title")
    existing_mode = spec.get("existing_mode", "insert")
    start = max(1, int(spec.get("start_slide", 1)))
    layout = _layout(prs, want_title=(title_mode == "title"))
    sld_list = prs.slides._sldIdLst
    insert_at = min(start - 1, len(prs.slides)) if mode == "existing" else len(prs.slides)

    count = 0
    for i, item in enumerate(spec["slides"]):
        new_slide = True
        if mode == "existing" and existing_mode == "reuse" and start - 1 + i < len(prs.slides):
            slide = prs.slides[start - 1 + i]
            new_slide = False
        else:
            slide = prs.slides.add_slide(layout)
            _remove_empty_placeholders(slide, keep_title=(title_mode == "title"))
            if mode == "existing" and existing_mode == "insert":
                el = sld_list[-1]
                sld_list.remove(el)
                sld_list.insert(insert_at, el)
                insert_at += 1
        _place_filename(slide, item.get("title", ""), title_mode, sw, sh, new_slide)
        top = int(sh * (0.145 if title_mode == "title" else 0.125 if title_mode == "subtitle" else 0.05))
        avail_h = int(sh - top - sh * 0.05)
        iw, ih = item.get("img_w", 0), item.get("img_h", 0)
        if item.get("marked"):
            text_left = int(sw * 0.04)
            text_w = int(sw * 0.36)
            gap = int(sw * 0.035)
            image_left = text_left + text_w + gap
            image_w = int(sw - image_left - sw * 0.04)
            pw, ph = _fit(iw, ih, image_w, avail_h)
            slide.shapes.add_picture(item["image"], image_left + (image_w - pw) // 2, top + (avail_h - ph) // 2, pw, ph)
            _add_comments(slide, item.get("comments") or [], text_left, top, text_w, avail_h)
        else:
            max_w = int(sw * 0.9)
            pw, ph = _fit(iw, ih, max_w, avail_h)
            slide.shapes.add_picture(item["image"], (sw - pw) // 2, top + max(avail_h - ph, 0) // 2, pw, ph)
        count += 1

    out = spec["output"]
    folder = os.path.dirname(os.path.abspath(out)) or "."
    fd, tmp = tempfile.mkstemp(prefix=".eng-export-", suffix=".pptx", dir=folder)
    os.close(fd)
    try:
        prs.save(tmp)
        try:
            os.replace(tmp, out)
        except PermissionError as exc:
            raise PermissionError(
                f"Could not write {os.path.basename(out)}. If it is open in PowerPoint, close it and try again."
            ) from exc
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return {"slides": count, "path": out, "total_slides": len(prs.slides)}
